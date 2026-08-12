"""
cli/curriculum_parser.py — Excel PM / Syllabus parsing and curriculum structure utilities.
"""

import re
from pathlib import Path
from datetime import datetime
import openpyxl

def sanitize_folder_name(name: str) -> str:
    """Loại bỏ các ký tự không hợp lệ cho tên thư mục trên mọi hệ điều hành và an toàn cho URL."""
    name = name.replace("&", "va").replace("%", "").replace("^", "").replace("#", "")
    sanitized = re.sub(r'[\\/*?:"<>|]', "", name).strip()
    # Loại bỏ dấu gạch nối/dấu gạch thừa ở cuối (ví dụ "Session 01 -" -> "Session 01")
    sanitized = re.sub(r'[\s\-_]+$', '', sanitized).strip()
    return sanitized

def format_full_folder_name(item_id: str, item_title: str) -> str:
    """Ghép ID và Title thư mục một cách thông minh, tránh lặp lại hoặc để lại gạch nối thừa ở cuối."""
    item_id = (item_id or "").strip()
    item_title = (item_title or "").strip()
    if not item_title or item_title == item_id or item_id.startswith(item_title):
        return item_id or item_title
    if item_title.startswith(item_id):
        return item_title
    return f"{item_id} - {item_title}"

import os

def _to_long_path(path: Path) -> Path:
    abs_path = os.path.abspath(str(path))
    if os.name == 'nt' and not abs_path.startswith('\\\\?\\'):
        abs_path = '\\\\?\\' + abs_path
    return Path(abs_path)

def get_or_rename_sanitized_folder(parent_dir: Path, prefix: str, full_name: str) -> Path:
    """
    Looks for an existing folder starting with prefix (e.g. 'Session 01').
    If found and it has a different name than the new sanitized full_name, renames it.
    This prevents duplicate session/lesson folders when titles change.
    """
    target_name = sanitize_folder_name(full_name)
    target_path = _to_long_path(parent_dir / target_name)
    parent_dir_long = _to_long_path(parent_dir)
    
    if parent_dir_long.exists():
        for item in parent_dir_long.iterdir():
            if item.is_dir():
                if item.name == target_name:
                    return target_path
                # Prevent matching "Session 010" when prefix is "Session 01"
                if item.name == prefix or item.name.startswith(prefix + " ") or item.name.startswith(prefix + "-"):
                    print(f"  [Auto-Rename] Đồng bộ thư mục cũ: '{item.name}' -> '{target_name}'")
                    try:
                        item.rename(target_path)
                    except Exception as e:
                        print(f"  [Warning] Không thể đổi tên thư mục cũ {item.name}: {e}")
                    return target_path
    return target_path

def parse_all_sessions(excel_path: str):
    """
    Parses all sessions and their corresponding lessons
    from the PM software engineering spreadsheet using dynamic column mapping.
    """
    wb = openpyxl.load_workbook(excel_path)
    
    target_sheet = "Chương trình đào tạo chi tiết"
    ws = wb[target_sheet] if target_sheet in wb.sheetnames else wb.active
    
    # Auto-detect PTIT detailed module sheet layout
    is_ptit_layout = False
    header_row = 3
    if ws.cell(row=3, column=1).value == "STT" or ws.cell(row=3, column=2).value == "Hình thức học":
        is_ptit_layout = True
        header_row = 3
    elif ws.cell(row=4, column=1).value == "STT" or ws.cell(row=4, column=2).value == "Hình thức học":
        is_ptit_layout = True
        header_row = 4
        
    if is_ptit_layout:
        sessions = []
        current_session = None
        rows = list(ws.iter_rows(values_only=True))
        
        # Start reading from row after header
        for r_idx in range(header_row, len(rows)):
            row = rows[r_idx]
            if not any(row):
                continue
                
            stt = row[0]
            hinh_thuc = row[1]
            session_title = row[4]
            lesson_idx = row[5]
            lesson_title = row[6]
            note = row[7] if len(row) > 7 else ""
            
            # If stt is present, it's a new Session
            if stt is not None and str(stt).strip() != "":
                stt_str = str(stt).strip()
                current_session = {
                    "session_id": f"Session {stt_str.zfill(2)}" if stt_str.isdigit() else f"Session {stt_str}",
                    "session_type": str(hinh_thuc).strip() if hinh_thuc else "Lý thuyết",
                    "session_code": f"SESS_{stt_str}",
                    "title": str(session_title).strip() if session_title else "",
                    "forbidden_scope": "",
                    "allowed_scope": "",
                    "tech_stack_convention": "",
                    "lessons": []
                }
                sessions.append(current_session)
                
            # If lesson_title is present and we have a session
            if lesson_title and str(lesson_title).strip() and current_session:
                l_id = f"Lesson {str(lesson_idx).strip()}" if lesson_idx else f"Lesson {len(current_session['lessons'])+1:02d}"
                current_session["lessons"].append({
                    "lesson_id": l_id,
                    "title": str(lesson_title).strip(),
                    "details": str(lesson_title).strip(),
                    "expected_output": "",
                    "forbidden_scope": "",
                    "allowed_scope": "",
                    "tech_stack_convention": ""
                })
        return sessions

    # 1. Identify header mapping dynamically for ALL 10 PM columns by scanning first 15 rows
    headers = {}
    header_row_idx = 1
    for r_idx, row_cells in enumerate(ws.iter_rows(min_row=1, max_row=15), start=1):
        for col_idx, cell in enumerate(row_cells, start=0):
            if cell.value and "session" in str(cell.value).strip().lower():
                header_row_idx = r_idx
                for c_idx, c_cell in enumerate(row_cells, start=0):
                    if c_cell.value:
                        val = str(c_cell.value).strip().lower()
                        if val == "session" or val == "stt session":
                            headers["session"] = c_idx
                        elif "loại session" in val or "loại" in val:
                            headers["session_type"] = c_idx
                        elif "mã session" in val:
                            headers["session_code"] = c_idx
                        elif "tên tiêu đề" in val or "tiêu đề session" in val or "chủ đề" in val:
                            headers["session_title"] = c_idx
                        elif "tên lesson" in val or val == "lesson" or val == "bài học":
                            headers["lesson"] = c_idx
                        elif "chi tiết" in val or "lesson scope" in val:
                            headers["details"] = c_idx
                        elif "kết quả mong đợi" in val or "expected" in val or "sản phẩm" in val:
                            headers["expected_output"] = c_idx
                        elif "cấm" in val or "forbidden" in val:
                            headers["forbidden_scope"] = c_idx
                        elif "đã học" in val or "allowed" in val:
                            headers["allowed_scope"] = c_idx
                        elif "tech stack" in val or "quy chuẩn" in val or "convention" in val:
                            headers["tech_stack_convention"] = c_idx
                break
        if headers:
            break

    # Fallback to hardcoded indexes matching PM_Python.xlsx standard columns [0..9]
    h_sess = headers.get("session", 0)
    h_s_type = headers.get("session_type", 1)
    h_s_code = headers.get("session_code", 2)
    h_s_title = headers.get("session_title", 3)
    h_lesson = headers.get("lesson", 4)
    h_details = headers.get("details", 5)
    h_out = headers.get("expected_output", 6)
    h_forbidden = headers.get("forbidden_scope", 7)
    h_allowed = headers.get("allowed_scope", 8)
    h_tech = headers.get("tech_stack_convention", 9)
    
    sessions = []
    current_session = None
    
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if not any(row): continue # Skip empty rows
        
        session_val = str(row[h_sess]).strip() if h_sess < len(row) and row[h_sess] else ""
        type_val = str(row[h_s_type]).strip() if h_s_type < len(row) and row[h_s_type] else ""
        code_val = str(row[h_s_code]).strip() if h_s_code < len(row) and row[h_s_code] else ""
        content_val = str(row[h_s_title]).strip() if h_s_title < len(row) and row[h_s_title] else ""
        lesson_val = str(row[h_lesson]).strip() if h_lesson < len(row) and row[h_lesson] else ""
        details_val = str(row[h_details]).strip() if h_details < len(row) and row[h_details] else ""
        output_val = str(row[h_out]).strip() if h_out < len(row) and row[h_out] else ""
        forbidden_val = str(row[h_forbidden]).strip() if h_forbidden < len(row) and row[h_forbidden] else ""
        allowed_val = str(row[h_allowed]).strip() if h_allowed < len(row) and row[h_allowed] else ""
        tech_val = str(row[h_tech]).strip() if h_tech < len(row) and row[h_tech] else ""
        
        # If a new Session is found (and not the table header row)
        if session_val.startswith("Session") and session_val.lower() not in ["session", "session / stt", "stt session"]:
            # Check if we already have it
            existing = [s for s in sessions if s["session_id"] == session_val]
            if existing:
                current_session = existing[0]
            else:
                current_session = {
                    "session_id": session_val,
                    "session_type": type_val,
                    "session_code": code_val,
                    "title": content_val,
                    "forbidden_scope": forbidden_val,
                    "allowed_scope": allowed_val,
                    "tech_stack_convention": tech_val,
                    "lessons": []
                }
                sessions.append(current_session)
            
        # Add lessons under the current session
        if lesson_val.startswith("Lesson") and current_session:
            # Try to safely split Lesson ID and Title
            if ":" in lesson_val:
                l_id, l_title = lesson_val.split(":", 1)
            elif "-" in lesson_val:
                l_id, l_title = lesson_val.split("-", 1)
            else:
                l_id, l_title = lesson_val, lesson_val
                
            current_session["lessons"].append({
                "lesson_id": l_id.strip(),
                "title": l_title.strip(),
                "details": details_val,
                "expected_output": output_val,
                "forbidden_scope": forbidden_val,
                "allowed_scope": allowed_val,
                "tech_stack_convention": tech_val,
                "session_type": type_val
            })
            
    return sessions

def get_context_curriculum(sessions: list, current_session_id: str) -> list:
    """
    To prevent LLM context window overflow, this function returns a truncated version
    of the curriculum containing only the previous, current, and next session.
    """
    idx = next((i for i, s in enumerate(sessions) if s["session_id"] == current_session_id), -1)
    if idx == -1: return sessions
    
    start = max(0, idx - 1)
    end = min(len(sessions), idx + 2)
    return sessions[start:end]


def is_exam_session(session: dict) -> bool:
    """Kiểm tra xem session có phải là dạng bài thi (Thi giữa kỳ, Thi thực hành, Thi cuối kỳ...)"""
    session_code = str(session.get("session_code", "")).strip().upper()
    session_title = str(session.get("title", "")).strip().lower()
    session_type = str(session.get("session_type", "")).strip().lower()

    if "ôn tập" in session_title or "ôn tập" in session_type:
        return False

    if session_code in ["MIDTERM", "FINAL", "EXAM"]:
        return True

    exam_keywords = ["thi giữa kỳ", "thi thực hành", "thi cuối kỳ", "thi giữa môn", "thi cuối môn", "midterm exam", "final exam"]
    if any(kw in session_type for kw in exam_keywords):
        return True

    # Avoid substring collision with 'thực hành thiết lập' by checking explicit exam title phrases
    if any(kw in session_title for kw in ["thi giữa kỳ", "thi cuối kỳ", "thi thực hành cuối kỳ", "thi thực hành giữa kỳ", "thi giữa môn", "thi cuối môn", "bài thi thực hành", "midterm exam", "final exam"]):
        return True

    return False


def is_project_or_hackathon_session(session: dict) -> bool:
    """Kiểm tra xem session có phải là dạng Mini Project / Capstone Project / Hackathon hay không"""
    session_code = str(session.get("session_code", "")).strip().upper()
    session_title = str(session.get("title", "")).strip().lower()
    session_type = str(session.get("session_type", "")).strip().lower()

    if session_code in ["MINI_PROJECT", "FINAL_PROJECT", "PROJECT", "HACKATHON"]:
        return True

    if session_type in ["mini project", "dự án cuối khóa", "đồ án", "hackathon"]:
        return True

    project_keywords = ["mini project", "hackathon", "dự án cuối khóa", "đồ án môn học", "đồ án tốt nghiệp", "capstone project"]
    if any(kw in session_title for kw in project_keywords) or any(kw in session_type for kw in project_keywords):
        return True

    return False


def is_practice_session(session: dict) -> bool:
    """Kiểm tra xem session có phải là dạng Thực hành (Practice Lab) hay không"""
    if is_exam_session(session) or is_project_or_hackathon_session(session):
        return False

    session_code = str(session.get("session_code", "")).strip().upper()
    session_title = str(session.get("title", "")).strip().lower()
    session_type = str(session.get("session_type", "")).strip().lower()

    if session_code == "PRACTICE":
        return True

    if "thực hành" in session_type or "thực hành" in session_title:
        return True

    return False


def initialize_skeleton_structure(sessions, course_dir: Path, requested_parts: list, requested_session: str):
    print("\n=====================================================================")
    print(">>> GIAI ĐOẠN 1.1: KHỞI TẠO KHUNG CẤU TRÚC THƯ MỤC & FILE RỖNG <<<")
    print("=====================================================================")
    for session in sessions:
        session_id = session["session_id"]
        session_title = session.get("title", "")
        req_sess_list = [r.strip().lower() for r in requested_session.split(",")] if requested_session != "all" else ["all"]
        if "all" not in req_sess_list and not any(rs in session_id.lower() for rs in req_sess_list):
            continue

        session_folder_name = format_full_folder_name(session_id, session_title)
        session_dir = get_or_rename_sanitized_folder(course_dir, session_id, session_folder_name)
        session_dir.mkdir(parents=True, exist_ok=True)
        print(f"  [Folder] Khởi tạo session: {session_folder_name}")

        is_exam = is_exam_session(session)
        is_project_or_hackathon = is_project_or_hackathon_session(session)
        is_practice = is_practice_session(session)

        if is_exam:
            (session_dir / "Đề thi thực hành").mkdir(parents=True, exist_ok=True)
            (session_dir / "Đề thi trắc nghiệm").mkdir(parents=True, exist_ok=True)
            (session_dir / "Câu hỏi vấn đáp").mkdir(parents=True, exist_ok=True)
            continue
        elif is_project_or_hackathon:
            (session_dir / "Bài kiểm tra đầu giờ").mkdir(parents=True, exist_ok=True)
            (session_dir / "Tài liệu đặc tả SRS").mkdir(parents=True, exist_ok=True)
            (session_dir / "Mini project").mkdir(parents=True, exist_ok=True)
            continue
        elif is_practice:
            (session_dir / "Bài tập").mkdir(parents=True, exist_ok=True)
            continue

        if session["lessons"]:
            for lesson in session["lessons"]:
                lesson_id = lesson["lesson_id"]
                lesson_title = lesson["title"]
                
                lesson_folder_name = format_full_folder_name(lesson_id, lesson_title)
                lesson_dir = get_or_rename_sanitized_folder(session_dir, lesson_id, lesson_folder_name)
                lesson_dir.mkdir(parents=True, exist_ok=True)

                is_lesson_project_or_hackathon = any(kw in lesson_title.lower() or kw in session_title.lower() for kw in ["hackathon", "project", "đồ án", "dự án", "mini project"])
                is_lesson_practice = "thực hành" in lesson_title.lower() or "thực hành" in session_title.lower()

                if is_lesson_project_or_hackathon:
                    (lesson_dir / "Bài kiểm tra đầu giờ").mkdir(parents=True, exist_ok=True)
                    (lesson_dir / "Tài liệu đặc tả SRS").mkdir(parents=True, exist_ok=True)
                    (lesson_dir / "Mini project").mkdir(parents=True, exist_ok=True)
                    continue
                elif is_lesson_practice:
                    (lesson_dir / "Bài tập").mkdir(parents=True, exist_ok=True)
                    continue
                
                if "html" in requested_parts:
                    sub = lesson_dir / "Bài đọc"
                    sub.mkdir(parents=True, exist_ok=True)
                    with open(sub / "reading.html", "w", encoding="utf-8") as f:
                        f.write(f"<!-- Empty outline for {session_id} - {lesson_id}: {lesson_title} -->\n")
                    
                    sub_q = lesson_dir / "Câu hỏi bài đọc"
                    sub_q.mkdir(parents=True, exist_ok=True)
                    with open(sub_q / "reading_questions.json", "w", encoding="utf-8") as f:
                        f.write("[]\n")
                
                # Note: Lesson-level Slides removed (Moved to Session level)
                
                if "quiz" in requested_parts:
                    sub = lesson_dir / "Câu hỏi Quizz"
                    sub.mkdir(parents=True, exist_ok=True)
                    with open(sub / "quiz.json", "w", encoding="utf-8") as f:
                        f.write("[]\n")
                    sub_lab = lesson_dir / "Bài thực hành"
                    sub_lab.mkdir(parents=True, exist_ok=True)
                    with open(sub_lab / "practical_lab.json", "w", encoding="utf-8") as f:
                        f.write("{}\n")
                
                # TẠM THỜI COMMENT LUỒNG TẠO SCRIPT VIDEO LESSON
                # if "video" in requested_parts or "video_script" in requested_parts:
                #     sub = lesson_dir / "Video"
                #     sub.mkdir(parents=True, exist_ok=True)
                #     with open(sub / "SCRIPT.md", "w", encoding="utf-8") as f:
                #         f.write(f"<!-- Empty video script outline for {session_id} - {lesson_id}: {lesson_title} -->\n")
                
                # Note: Lesson-level Mindmap removed (Moved to Session level)
        else:
            if "html" in requested_parts:
                sub = session_dir / "Bài đọc"
                sub.mkdir(parents=True, exist_ok=True)
                with open(sub / "reading.html", "w", encoding="utf-8") as f:
                    f.write(f"<!-- Empty outline for {session_id} -->\n")
            # TẠM THỜI COMMENT LUỒNG TẠO BÀI GIẢNG SLIDE SESSION
            # if "slide" in requested_parts:
            #     sub = session_dir / "Bài giảng"
            #     sub.mkdir(parents=True, exist_ok=True)
            #     with open(sub / "slides.html", "w", encoding="utf-8") as f:
            #         f.write(f"<!-- Empty slide outline for {session_id} -->\n")
            if "quiz" in requested_parts:
                sub = session_dir / "Câu hỏi Quizz"
                sub.mkdir(parents=True, exist_ok=True)
                with open(sub / "quiz.json", "w", encoding="utf-8") as f:
                    f.write("[]\n")
                sub_lab = session_dir / "Bài thực hành"
                sub_lab.mkdir(parents=True, exist_ok=True)
                with open(sub_lab / "practical_lab.json", "w", encoding="utf-8") as f:
                    f.write("{}\n")
            # TẠM THỜI COMMENT LUỒNG TẠO SCRIPT VIDEO SESSION
            # if "video" in requested_parts or "video_script" in requested_parts:
            #     sub = session_dir / "Video"
            #     sub.mkdir(parents=True, exist_ok=True)
            #     with open(sub / "SCRIPT.md", "w", encoding="utf-8") as f:
            #         f.write(f"<!-- Empty video script outline for {session_id} -->\n")
            if "mindmap" in requested_parts:
                sub = session_dir / "Mindmap"
                sub.mkdir(parents=True, exist_ok=True)
                with open(sub / "mindmap.md", "w", encoding="utf-8") as f:
                    f.write(f"<!-- Empty mindmap outline for {session_id} -->\n")

def project_structure_reviewer_agent(sessions, course_dir: Path, requested_parts: list, requested_session: str):
    print("\n=====================================================================")
    print(">>> GIAI ĐOẠN 1.2: AGENT QUÉT & THẨM ĐỊNH CẤU TRÚC KHUNG HỌC LIỆU <<<")
    print("=====================================================================")
    missing_elements = []
    for session in sessions:
        session_id = session["session_id"]
        session_title = session.get("title", "")
        req_sess_list = [r.strip() for r in requested_session.split(",")] if requested_session != "all" else ["all"]
        if "all" not in req_sess_list and not any(rs in session_id.lower() for rs in req_sess_list):
            continue

        session_dir = get_or_rename_sanitized_folder(course_dir, session_id, format_full_folder_name(session_id, session_title))
        if not session_dir.exists():
            missing_elements.append(f"Thiếu thư mục session: {session_id}")
            continue

        is_exam = is_exam_session(session)
        is_project_or_hackathon = is_project_or_hackathon_session(session)
        is_practice = is_practice_session(session)

        if is_exam:
            if not (session_dir / "Đề thi thực hành").exists():
                missing_elements.append(f"Thiếu thư mục 'Đề thi thực hành' tại {session_id}")
            if not (session_dir / "Đề thi trắc nghiệm").exists():
                missing_elements.append(f"Thiếu thư mục 'Đề thi trắc nghiệm' tại {session_id}")
            if not (session_dir / "Câu hỏi vấn đáp").exists():
                missing_elements.append(f"Thiếu thư mục 'Câu hỏi vấn đáp' tại {session_id}")
            continue
        elif is_project_or_hackathon:
            if not (session_dir / "Bài kiểm tra đầu giờ").exists():
                missing_elements.append(f"Thiếu thư mục 'Bài kiểm tra đầu giờ' tại {session_id}")
            if not (session_dir / "Tài liệu đặc tả SRS").exists():
                missing_elements.append(f"Thiếu thư mục 'Tài liệu đặc tả SRS' tại {session_id}")
            if not (session_dir / "Mini project").exists():
                missing_elements.append(f"Thiếu thư mục 'Mini project' tại {session_id}")
            continue
        elif is_practice:
            if not (session_dir / "Bài tập").exists():
                missing_elements.append(f"Thiếu thư mục 'Bài tập' tại {session_id}")
            continue

        if session["lessons"]:
            for lesson in session["lessons"]:
                lesson_id = lesson["lesson_id"]
                lesson_title = lesson["title"]
                lesson_dir = get_or_rename_sanitized_folder(session_dir, lesson_id, format_full_folder_name(lesson_id, lesson_title))
                if not lesson_dir.exists():
                    missing_elements.append(f"Thiếu thư mục lesson: {session_id} -> {lesson_id}")
                    continue

                is_lesson_project_or_hackathon = any(kw in lesson_title.lower() or kw in session_title.lower() for kw in ["hackathon", "project", "đồ án", "dự án", "mini project"])
                is_lesson_practice = "thực hành" in lesson_title.lower() or "thực hành" in session_title.lower()

                if is_lesson_project_or_hackathon:
                    if not (lesson_dir / "Bài kiểm tra đầu giờ").exists():
                        missing_elements.append(f"Thiếu thư mục 'Bài kiểm tra đầu giờ' tại {session_id} -> {lesson_id}")
                    if not (lesson_dir / "Tài liệu đặc tả SRS").exists():
                        missing_elements.append(f"Thiếu thư mục 'Tài liệu đặc tả SRS' tại {session_id} -> {lesson_id}")
                    if not (lesson_dir / "Mini project").exists():
                        missing_elements.append(f"Thiếu thư mục 'Mini project' tại {session_id} -> {lesson_id}")
                    continue
                elif is_lesson_practice:
                    if not (lesson_dir / "Bài tập").exists():
                        missing_elements.append(f"Thiếu thư mục 'Bài tập' tại {session_id} -> {lesson_id}")
                    continue

                if "html" in requested_parts and not (lesson_dir / "Bài đọc" / "reading.html").exists():
                    missing_elements.append(f"Thiếu file reading.html tại {session_id} -> {lesson_id}")
                if "quiz" in requested_parts:
                    if not (lesson_dir / "Câu hỏi Quizz" / "quiz.json").exists():
                        missing_elements.append(f"Thiếu file quiz.json tại {session_id} -> {lesson_id}")
                    if not (lesson_dir / "Bài thực hành" / "practical_lab.json").exists():
                        missing_elements.append(f"Thiếu file practical_lab.json tại {session_id} -> {lesson_id}")
                # TẠM THỜI COMMENT KIỂM TRA SCRIPT VIDEO LESSON
                # if ("video" in requested_parts or "video_script" in requested_parts) and not (lesson_dir / "Video" / "SCRIPT.md").exists():
                #     missing_elements.append(f"Thiếu file SCRIPT.md tại {session_id} -> {lesson_id}")
        else:
            if "html" in requested_parts and not (session_dir / "Bài đọc" / "reading.html").exists():
                missing_elements.append(f"Thiếu file reading.html tại {session_id}")
            # TẠM THỜI COMMENT KIỂM TRA BÀI GIẢNG SLIDE VÀ SCRIPT VIDEO SESSION
            # if "slide" in requested_parts and not (session_dir / "Bài giảng" / "slides.html").exists():
            #     missing_elements.append(f"Thiếu file slides.html tại {session_id}")
            if "quiz" in requested_parts:
                if not (session_dir / "Câu hỏi Quizz" / "quiz.json").exists():
                    missing_elements.append(f"Thiếu file quiz.json tại {session_id}")
                if not (session_dir / "Bài thực hành" / "practical_lab.json").exists():
                    missing_elements.append(f"Thiếu file practical_lab.json tại {session_id}")
            # if ("video" in requested_parts or "video_script" in requested_parts) and not (session_dir / "Video" / "SCRIPT.md").exists():
            #     missing_elements.append(f"Thiếu file SCRIPT.md tại {session_id}")
            if "mindmap" in requested_parts and not (session_dir / "Mindmap" / "session_mindmap.md").exists() and not (session_dir / "Mindmap" / "mindmap.md").exists():
                missing_elements.append(f"Thiếu file session_mindmap.md tại {session_id}")

    report_path = course_dir / "structure_review_report.md"
    status = "APPROVED" if not missing_elements else "REJECTED"
    
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    report_content = f"""# 📂 BÁO CÁO DUYỆT KHUNG CẤU TRÚC DỰ ÁN (PROJECT STRUCTURE REVIEW)
**Đường dẫn dự án:** `{course_dir}`

## 📊 Kết quả Thẩm định
* **Trạng thái:** `{status}`
* **Thời gian quét:** {now_str}

## 🔍 Chi tiết đánh giá
{"Mọi thư mục Session và Lesson rỗng đã được tạo lập thành công và đầy đủ cấu trúc khung rỗng." if not missing_elements else "Phát hiện các lỗi cấu trúc sau:"}

{chr(10).join([f"* ❌ {err}" for err in missing_elements]) if missing_elements else "* ✅ Cấu trúc thư mục đạt chuẩn outline ban đầu.*"}
"""
    with open(report_path, "w", encoding="utf-8") as f:
        f.write(report_content)
    print(f"  [Structure Reviewer] Đã lưu báo cáo duyệt cấu trúc tại: {report_path.resolve()}")

    if missing_elements:
        raise ValueError(f"Duyệt cấu trúc thất bại. Xem chi tiết tại {report_path.name}")
    else:
        print("  [Structure Reviewer] ĐÃ PHÊ DUYỆT cấu trúc dự án. Tiếp tục giai đoạn xây dựng nội dung học liệu...")

def verify_previous_lessons_completed(sessions, current_session_id: str, current_lesson_id: str, course_dir: Path, requested_parts: list):
    # Flatten all lessons across sessions in sequence
    ordered_lessons = []
    for session in sessions:
        s_id = session["session_id"]
        s_title = session["title"]
        if session["lessons"]:
            for lesson in session["lessons"]:
                l_id = lesson["lesson_id"]
                l_title = lesson["title"]
                ordered_lessons.append((s_id, s_title, l_id, l_title))
        else:
            ordered_lessons.append((s_id, s_title, "", ""))
            
    # Find index of current
    current_idx = -1
    for idx, (s_id, _, l_id, _) in enumerate(ordered_lessons):
        if s_id == current_session_id and (l_id == current_lesson_id or (not l_id and not current_lesson_id)):
            current_idx = idx
            break
            
    if current_idx <= 0:
        return
        
    # Check all lessons before current_idx
    for idx in range(current_idx):
        prev_session_id, prev_session_title, prev_lesson_id, prev_lesson_title = ordered_lessons[idx]
        
        is_project_or_hackathon = any(kw in prev_session_title.lower() or kw in prev_lesson_title.lower() for kw in ["hackathon", "project", "đồ án", "dự án", "mini project"])
        is_practice = "thực hành" in prev_session_title.lower() or "thực hành" in prev_lesson_title.lower()
        if is_project_or_hackathon or is_practice:
            continue
            
        prev_session_dir = get_or_rename_sanitized_folder(course_dir, prev_session_id, format_full_folder_name(prev_session_id, prev_session_title))
        
        if prev_lesson_id:
            prev_dir = get_or_rename_sanitized_folder(prev_session_dir, prev_lesson_id, format_full_folder_name(prev_lesson_id, prev_lesson_title))
            name_str = f"{prev_session_id} -> {prev_lesson_id}"
        else:
            prev_dir = prev_session_dir
            name_str = prev_session_id
            
        uncompleted = []
        if "html" in requested_parts:
            html_file = prev_dir / "Bài đọc" / "reading.html"
            if not html_file.exists() or html_file.stat().st_size < 500 or "<!-- Empty outline" in html_file.read_text(encoding="utf-8"):
                uncompleted.append("Bài đọc/reading.html chưa được sinh nội dung chi tiết hoặc quá ngắn.")
        if "quiz" in requested_parts:
            quiz_file = prev_dir / "Câu hỏi Quizz" / "quiz.json"
            if not quiz_file.exists() or quiz_file.stat().st_size < 20:
                uncompleted.append("Câu hỏi Quizz/quiz.json chưa được sinh nội dung chi tiết.")
            
            lab_file = prev_dir / "Bài thực hành" / "practical_lab.json"
            if not lab_file.exists() or lab_file.stat().st_size < 20:
                uncompleted.append("Bài thực hành/practical_lab.json chưa được sinh nội dung chi tiết.")
        if "video" in requested_parts or "video_script" in requested_parts:
            video_file = prev_dir / "Video" / "SCRIPT.md"
            if not video_file.exists() or video_file.stat().st_size < 500 or "<!-- Empty video script outline" in video_file.read_text(encoding="utf-8"):
                uncompleted.append("Video/SCRIPT.md chưa được sinh nội dung chi tiết hoặc quá ngắn.")
                
        if uncompleted:
            feedback_details = "\n".join([f"  - {item}" for item in uncompleted])
            raise ValueError(
                f"\n[RÀNG BUỘC TUẦN TỰ NGHIÊM NGẶT] Bắt buộc hoàn thành và phê duyệt học liệu trước đó:\n"
                f"📌 {name_str} chưa được sinh hoặc chưa đạt chuẩn phê duyệt:\n"
                f"{feedback_details}\n"
                f"Yêu cầu: AI phải tuân thủ kỷ luật sư phạm, thà làm chậm và chất lượng còn hơn làm ẩu làm nhanh!"
            )

def parse_program_structure_from_ptit(excel_path: str, target_course_id: str):
    """
    Scans the 5 semester overview sheets to find target_course_id,
    extracts its course metadata, PLOs, and CLOs.
    """
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # We will search the first 5 sheets (Semesters) for target_course_id
    semester_sheets = [s for s in wb.sheetnames if "Khung chương trình" in s]
    
    course_info = None
    
    for sheetname in semester_sheets:
        ws = wb[sheetname]
        rows = list(ws.iter_rows(values_only=True))
        
        # We need to find PLOs first (usually in row 3, index 2)
        # Let's search all rows for "PLO" title
        plos = []
        for r_idx, row in enumerate(rows[:5]):
            if row[0] and "plo" in str(row[0]).lower():
                # E3 is at column index 4
                plo_val = row[4] if len(row) > 4 else ""
                if plo_val:
                    plos = [p.strip() for p in str(plo_val).split("\n") if p.strip()]
                break
                
        # Now find the course row
        # Col 1 is "Mã môn". Let's scan all rows starting from row index 4 (5th row)
        current_semester_val = ""
        for r_idx in range(4, len(rows)):
            row = rows[r_idx]
            if not any(row):
                continue
            
            # Fill forward for semester (Col 2)
            if len(row) > 2 and row[2] is not None:
                current_semester_val = str(row[2]).strip()
                
            course_id_cell = row[1] if len(row) > 1 else None
            if course_id_cell and str(course_id_cell).strip().lower() == target_course_id.lower().strip():
                # Found the course!
                course_name = str(row[3]).strip().replace("\n", " ") if len(row) > 3 and row[3] else ""
                clos_raw = str(row[4]).strip() if len(row) > 4 and row[4] else ""
                clos = [c.strip() for c in clos_raw.split("\n") if c.strip()]
                
                main_content = str(row[5]).strip() if len(row) > 5 and row[5] else ""
                exam_type = str(row[6]).strip() if len(row) > 6 and row[6] else ""
                
                lt = int(row[7]) if len(row) > 7 and row[7] is not None else 0
                th_off = int(row[8]) if len(row) > 8 and row[8] is not None else 0
                th_on = int(row[9]) if len(row) > 9 and row[9] is not None else 0
                mini_proj = int(row[10]) if len(row) > 10 and row[10] is not None else 0
                exam = int(row[11]) if len(row) > 11 and row[11] is not None else 0
                proj = int(row[12]) if len(row) > 12 and row[12] is not None else 0
                
                course_info = {
                    "course_id": target_course_id.upper().strip(),
                    "course_name": course_name,
                    "semester_id": current_semester_val,
                    "semester_sheet": sheetname,
                    "plos": plos,
                    "clos": clos,
                    "main_content": main_content,
                    "exam_type": exam_type,
                    "hours": {
                        "theory": lt,
                        "practice_offline": th_off,
                        "practice_online": th_on,
                        "practice_total": th_off + th_on,
                        "mini_project": mini_proj,
                        "exam": exam,
                        "project": proj,
                        "total": lt + th_off + th_on + mini_proj + exam + proj,
                        "hours_per_session": 2.0
                    }
                }
                break
        if course_info:
            break
            
    return course_info


def detect_tech_stack_from_curriculum(sessions: list, excel_path: str) -> str:
    """
    Dynamically extracts technology stack and pedagogical domain for ANY course 
    by reading Excel metadata or querying the LLM Router.
    ABSOLUTELY NO hardcoded course names, keywords, or language arrays.
    Works dynamically for any subject (Git, Agile/Scrum, System Analysis & Design, AI Tools, Testing, UI/UX, etc.).
    """
    if not sessions:
        return "general/software_engineering"

    # 1. Read explicit tech_stack_convention from Excel PM metadata if provided by author
    for sess in sessions:
        if sess.get("tech_stack_convention"):
            return sess["tech_stack_convention"].strip().lower()

    # 2. Extract curriculum outline summary for dynamic LLM domain classification
    course_name = Path(excel_path).stem.replace("PM_Generated_", "").replace("_Updated", "").replace("_", " ")
    session_titles = [s.get("title", "") for s in sessions[:5] if s.get("title")]
    outline_summary = f"Môn học: {course_name}\nCác session tiêu biểu: " + ", ".join(session_titles)

    try:
        from core.llm_router import call_llm
        prompt = f"""Bạn là Chuyên gia Kiến trúc Sư phạm (Curriculum Domain Architect).
Hãy phân tích thông tin môn học và xuất ra ĐÚNG 1 TỪ KHÓA ĐỊNH DẠNG DOMAIN/TECH-STACK ngắn gọn đại diện cho công nghệ/miền tri thức chính của môn học này.

Ví dụ định dạng trả về (chỉ trả về 1 chuỗi dạng category/domain, không giải thích):
- Môn Quản lý phiên bản Git -> vcs/git
- Môn Phân tích thiết kế hệ thống -> system_analysis/uml
- Môn Phương pháp Agile Scrum -> agile_scrum
- Môn Ứng dụng AI trong học tập -> ai_tools/prompting
- Môn Lập trình Python -> python/core
- Môn Cơ sở dữ liệu SQL -> database/sql
- Môn Lập trình Web Frontend -> web/frontend
- Môn Kiểm thử phần mềm -> software_testing/qa

Thông tin môn học:
{outline_summary}

Chỉ trả về 1 chuỗi domain/tech-stack dạng category/technology ngắn gọn:"""
        
        domain = call_llm(prompt, temperature=0.0).strip().lower()
        domain = domain.replace("`", "").replace('"', '').replace("'", "").strip()
        if domain and len(domain) < 50:
            return domain
    except Exception as e:
        print(f"  [Warning] Dynamic LLM domain detection fallback: {e}")

    return "general/software_engineering"
