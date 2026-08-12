try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False
    Presentation = None
    RGBColor = lambda *a: None
    Inches = Pt = lambda *a: None
    class _DummyEnum:
        def __getattr__(self, name): return None
    PP_ALIGN = MSO_ANCHOR = MSO_SHAPE = _DummyEnum()

class PPTXGeneratorAgent:
    """
    Presenton-inspired Native PowerPoint (.pptx) Generator Engine for Rikkei Elearning Agent.
    Compiles session data & slide scenes into a downloadable 16:9 widescreen PowerPoint presentation (.pptx).
    """

    # Brand Palette
    COLOR_RIKKEI_RED = RGBColor(190, 17, 28)       # #be111c
    COLOR_DARK_SLATE = RGBColor(15, 23, 42)       # #0f172a
    COLOR_CODE_BG = RGBColor(30, 41, 59)          # #1e293b
    COLOR_CARD_RED_BG = RGBColor(254, 242, 242)    # #fef2f2
    COLOR_CARD_RED_BORDER = RGBColor(254, 202, 202)# #fecaca
    COLOR_CARD_EMERALD_BG = RGBColor(236, 253, 245)# #ecfdf5
    COLOR_CARD_EMERALD_BORDER = RGBColor(167, 243, 208) # #a7f3d0
    COLOR_CARD_BLUE_BG = RGBColor(239, 246, 255)   # #eff6ff
    COLOR_CARD_BLUE_BORDER = RGBColor(191, 219, 254)  # #bfdbfe
    COLOR_CARD_SLATE_BG = RGBColor(248, 250, 252)  # #f8fafc
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_TEXT_MAIN = RGBColor(15, 23, 42)
    COLOR_TEXT_MUTED = RGBColor(100, 116, 139)

    FONT_TITLE = "Montserrat"
    FONT_BODY = "Arial"
    FONT_CODE = "Consolas"

    def clean_title_string(self, title: str) -> str:
        if not title:
            return ""
        clean = re.sub(r'</?[^>]+>', '', title)
        clean = re.sub(r'^\s*(Session\s*\d+|Lesson\s*\d+|Bài\s*\d+|\[\d+\.\d+\]|\d+\.)\s*[\:\-]?\s*', '', clean, flags=re.IGNORECASE).strip()
        clean = re.sub(r'^(Session\s*\d+\s*[\:\-]?\s*)+', '', clean, flags=re.IGNORECASE).strip()
        return clean

    def _create_base_presentation(self) -> Presentation:
        prs = Presentation()
        # Set 16:9 Widescreen Aspect Ratio (13.333" x 7.5")
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        return prs

    def _add_footer(self, slide, page_num: int, total_pages: int, copyright_text: str):
        # Footer text
        txBox = slide.shapes.add_textbox(Inches(3.0), Inches(7.0), Inches(7.333), Inches(0.4))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = copyright_text
        p.font.name = self.FONT_BODY
        p.font.size = Pt(11)
        p.font.color.rgb = self.COLOR_TEXT_MUTED
        p.alignment = PP_ALIGN.CENTER

        # Red triangle badge at bottom-right
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_TRIANGLE,
            Inches(12.333), Inches(6.5), Inches(1.0), Inches(1.0)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.COLOR_RIKKEI_RED
        shape.line.fill.background()
        shape.rotation = 180

        # Page number text inside badge
        num_box = slide.shapes.add_textbox(Inches(12.5), Inches(6.7), Inches(0.7), Inches(0.6))
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = str(page_num)
        p_num.font.name = self.FONT_TITLE
        p_num.font.size = Pt(14)
        p_num.font.bold = True
        p_num.font.color.rgb = self.COLOR_WHITE
        p_num.alignment = PP_ALIGN.CENTER

    def _add_header(self, slide, title_text: str, subtitle_text: str):
        # Red main title
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = self.FONT_TITLE
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = self.COLOR_RIKKEI_RED

        # Subtitle
        if subtitle_text:
            tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.733), Inches(0.4))
            tf_sub = tb_sub.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle_text
            p_sub.font.name = self.FONT_BODY
            p_sub.font.size = Pt(15)
            p_sub.font.bold = True
            p_sub.font.color.rgb = self.COLOR_TEXT_MAIN

    def generate_deck_pptx(
        self,
        session_title: str,
        module_name: str,
        lessons_data: List[Dict[str, Any]],
        output_path: Optional[str] = None,
        core_ssot: Optional[Dict[str, Any]] = None
    ) -> bytes:
        """
        Generates a 16:9 Widescreen PowerPoint Presentation (.pptx).
        """
        if not HAS_PPTX:
            print("  [PPTX Generator Warning] python-pptx is not installed. Skipping PPTX generation.")
            return b""
        prs = self._create_base_presentation()
        blank_layout = prs.slide_layouts[6]
        copyright_text = "© 2026 By Rikkei Education - All rights reserved."

        clean_session = self.clean_title_string(session_title)
        m_sess = re.search(r'^(Session\s*\d+)\s*[:-]?\s*(.*)', session_title, re.IGNORECASE)
        session_tag = m_sess.group(1).strip() if m_sess else "Session 01"
        main_title = m_sess.group(2).strip() if m_sess else clean_session

        # ── SLIDE 1: COVER SLIDE ────────────────────────────────────────────
        slide0 = prs.slides.add_slide(blank_layout)

        # Left red polygon banner shape
        left_banner = slide0.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(2.2), Inches(0.3), Inches(3.0))
        left_banner.fill.solid()
        left_banner.fill.fore_color.rgb = self.COLOR_RIKKEI_RED
        left_banner.line.fill.background()

        # Session Tag
        tb_tag = slide0.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.5), Inches(0.6))
        p_tag = tb_tag.text_frame.paragraphs[0]
        p_tag.text = f"{session_tag}:"
        p_tag.font.name = self.FONT_TITLE
        p_tag.font.size = Pt(28)
        p_tag.font.bold = True
        p_tag.font.color.rgb = self.COLOR_RIKKEI_RED

        # Main Title
        tb_title = slide0.shapes.add_textbox(Inches(1.2), Inches(2.9), Inches(10.5), Inches(1.5))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_main = tf_title.paragraphs[0]
        p_main.text = main_title
        p_main.font.name = self.FONT_TITLE
        p_main.font.size = Pt(36)
        p_main.font.bold = True
        p_main.font.color.rgb = self.COLOR_TEXT_MAIN

        # Module / Course Name
        tb_mod = slide0.shapes.add_textbox(Inches(1.2), Inches(4.5), Inches(10.5), Inches(0.5))
        p_mod = tb_mod.text_frame.paragraphs[0]
        p_mod.text = f"Môn học: {module_name}"
        p_mod.font.name = self.FONT_BODY
        p_mod.font.size = Pt(16)
        p_mod.font.color.rgb = self.COLOR_TEXT_MUTED

        self._add_footer(slide0, 1, 10, copyright_text)

        # ── SLIDE 2: AGENDA SLIDE ───────────────────────────────────────────
        slide1 = prs.slides.add_slide(blank_layout)
        tb_ag_header = slide1.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.733), Inches(0.6))
        p_ag = tb_ag_header.text_frame.paragraphs[0]
        p_ag.text = "NỘI DUNG BÀI HỌC"
        p_ag.font.name = self.FONT_TITLE
        p_ag.font.size = Pt(28)
        p_ag.font.bold = True
        p_ag.font.color.rgb = self.COLOR_RIKKEI_RED

        tb_list = slide1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
        tf_list = tb_list.text_frame
        tf_list.word_wrap = True

        for idx, l_data in enumerate(lessons_data, 1):
            l_title = l_data.get("lesson_title") or f"Bài học {idx}"
            clean_lt = self.clean_title_string(l_title)
            clean_lt = re.sub(r'^\s*(Session\s*\d+|Lesson\s*\d+|Bài\s*\d+|\[\d+\.\d+\]|\d+\.)\s*[\:\-]?\s*', '', clean_lt, flags=re.IGNORECASE).strip()
            clean_lt = re.sub(r'^(Session\s*\d+\s*[\:\-]?\s*)+', '', clean_lt, flags=re.IGNORECASE).strip()
            p_item = tf_list.add_paragraph() if idx > 1 else tf_list.paragraphs[0]
            p_item.text = f"{idx}.  {clean_lt}"
            p_item.font.name = self.FONT_BODY
            p_item.font.size = Pt(22)
            p_item.font.bold = True
            p_item.font.color.rgb = self.COLOR_TEXT_MAIN
            p_item.space_after = Pt(22)

        self._add_footer(slide1, 2, 10, copyright_text)

        # ── SLIDE 3..N: CONTENT SLIDES ───────────────────────────────────────
        page_counter = 3

        for l_idx, l_data in enumerate(lessons_data, 1):
            l_title = l_data.get("lesson_title") or f"Bài học {l_idx}"
            clean_lt = self.clean_title_string(l_title)
            scenes = l_data.get("scenes", [])

            for s_idx, scene in enumerate(scenes, 1):
                slide_c = prs.slides.add_slide(blank_layout)
                stitle = scene.get("action_title") or scene.get("short_title") or scene.get("scene_title") or f"Chủ đề {s_idx}"
                clean_st = self.clean_title_string(stitle)

                header_main = f"{l_idx}. {clean_lt} - {s_idx}"
                self._add_header(slide_c, header_main, clean_st)

                # Layout logic
                html_raw = scene.get("html_content") or ""
                code_sample = scene.get("code_sample") or scene.get("code") or ""
                mermaid_code = scene.get("mermaid") or scene.get("diagram") or ""

                if "bento-card" in html_raw and "bg-red-500" in html_raw:
                    # 2-Column Bento Card Layout (Red vs Emerald)
                    card1 = slide_c.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.7))
                    card1.fill.solid()
                    card1.fill.fore_color.rgb = self.COLOR_CARD_RED_BG
                    card1.line.color.rgb = self.COLOR_CARD_RED_BORDER

                    tf1 = card1.text_frame
                    tf1.word_wrap = True
                    p1 = tf1.paragraphs[0]
                    p1.text = "Thách thức & Hạn chế phương pháp cũ"
                    p1.font.name = self.FONT_BODY
                    p1.font.size = Pt(16)
                    p1.font.bold = True
                    p1.font.color.rgb = RGBColor(153, 27, 27)

                    p1_desc = tf1.add_paragraph()
                    p1_desc.text = "\n• Dễ ghi đè mã nguồn khi làm việc nhóm thủ công.\n• Khó truy vết lịch sử chỉnh sửa và tác giả phát sinh lỗi."
                    p1_desc.font.name = self.FONT_BODY
                    p1_desc.font.size = Pt(14)
                    p1_desc.font.color.rgb = self.COLOR_TEXT_MAIN

                    card2 = slide_c.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(4.7))
                    card2.fill.solid()
                    card2.fill.fore_color.rgb = self.COLOR_CARD_EMERALD_BG
                    card2.line.color.rgb = self.COLOR_CARD_EMERALD_BORDER

                    tf2 = card2.text_frame
                    tf2.word_wrap = True
                    p2 = tf2.paragraphs[0]
                    p2.text = "Giải pháp hiện đại chuẩn doanh nghiệp"
                    p2.font.name = self.FONT_BODY
                    p2.font.size = Pt(16)
                    p2.font.bold = True
                    p2.font.color.rgb = RGBColor(6, 95, 70)

                    p2_desc = tf2.add_paragraph()
                    p2_desc.text = "\n• Tự động hóa quản lý các phiên bản mã nguồn.\n• Hỗ trợ làm việc nhóm đa chi nhánh an toàn 100%."
                    p2_desc.font.name = self.FONT_BODY
                    p2_desc.font.size = Pt(14)
                    p2_desc.font.color.rgb = self.COLOR_TEXT_MAIN

                elif mermaid_code or "MERMAID" in str(scene.get("layout_type")).upper():
                    # Mermaid Process Box
                    box = slide_c.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(4.7))
                    box.fill.solid()
                    box.fill.fore_color.rgb = self.COLOR_CARD_SLATE_BG
                    box.line.color.rgb = RGBColor(226, 232, 240)

                    tf_m = box.text_frame
                    tf_m.word_wrap = True
                    p_m_title = tf_m.paragraphs[0]
                    p_m_title.text = "SƠ ĐỒ LUỒNG VẬN HÀNH & QUY TRÌNH HỆ THỐNG"
                    p_m_title.font.name = self.FONT_TITLE
                    p_m_title.font.size = Pt(16)
                    p_m_title.font.bold = True
                    p_m_title.font.color.rgb = self.COLOR_RIKKEI_RED
                    p_m_title.alignment = PP_ALIGN.CENTER

                    p_m_code = tf_m.add_paragraph()
                    p_m_code.text = f"\n{mermaid_code.strip()}"
                    p_m_code.font.name = self.FONT_CODE
                    p_m_code.font.size = Pt(13)
                    p_m_code.font.color.rgb = self.COLOR_TEXT_MAIN

                else:
                    # Default Theory + Terminal Code Box
                    card_left = slide_c.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.7))
                    card_left.fill.solid()
                    card_left.fill.fore_color.rgb = self.COLOR_CARD_SLATE_BG
                    card_left.line.color.rgb = RGBColor(226, 232, 240)

                    tf_l = card_left.text_frame
                    tf_l.word_wrap = True
                    p_l_t = tf_l.paragraphs[0]
                    p_l_t.text = "Nguyên lý & Quy chuẩn thực thi"
                    p_l_t.font.name = self.FONT_BODY
                    p_l_t.font.size = Pt(16)
                    p_l_t.font.bold = True
                    p_l_t.font.color.rgb = self.COLOR_TEXT_MAIN

                    p_l_d = tf_l.add_paragraph()
                    p_l_d.text = "\n• Nắm vững các bước thao tác thực chiến.\n• Tuân thủ nguyên tắc an toàn và kiểm thử mã nguồn trước khi commit."
                    p_l_d.font.name = self.FONT_BODY
                    p_l_d.font.size = Pt(14)
                    p_l_d.font.color.rgb = self.COLOR_TEXT_MUTED

                    # Dark Slate Code Box
                    card_code = slide_c.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(4.7))
                    card_code.fill.solid()
                    card_code.fill.fore_color.rgb = self.COLOR_CODE_BG
                    card_code.line.fill.background()

                    tf_c = card_code.text_frame
                    tf_c.word_wrap = True
                    p_c = tf_c.paragraphs[0]
                    p_c.text = code_sample if code_sample else "# Thực thi lệnh thực chiến\n$ git status\n$ git add .\n$ git commit -m 'feat: update module'"
                    p_c.font.name = self.FONT_CODE
                    p_c.font.size = Pt(12)
                    p_c.font.color.rgb = RGBColor(241, 245, 249)

                self._add_footer(slide_c, page_counter, 10, copyright_text)
                page_counter += 1

        # ── SLIDE N: SUMMARY SLIDE ──────────────────────────────────────────
        slide_summary = prs.slides.add_slide(blank_layout)
        tb_sum_header = slide_summary.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.733), Inches(0.6))
        p_sum = tb_sum_header.text_frame.paragraphs[0]
        p_sum.text = "TỔNG KẾT BÀI HỌC"
        p_sum.font.name = self.FONT_TITLE
        p_sum.font.size = Pt(28)
        p_sum.font.bold = True
        p_sum.font.color.rgb = self.COLOR_RIKKEI_RED

        tb_sum_list = slide_summary.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
        tf_sum_list = tb_sum_list.text_frame
        tf_sum_list.word_wrap = True

        from agents.slide_generator_agent import slide_generator_agent
        summary_bullets = slide_generator_agent.extract_session_summary_bullets(lessons_data, core_ssot)

        for s_idx, s_bullet in enumerate(summary_bullets, 1):
            p_sitem = tf_sum_list.add_paragraph() if s_idx > 1 else tf_sum_list.paragraphs[0]
            p_sitem.text = f"✔  {s_bullet}"
            p_sitem.font.name = self.FONT_BODY
            p_sitem.font.size = Pt(20)
            p_sitem.font.bold = True
            p_sitem.font.color.rgb = self.COLOR_TEXT_MAIN
            p_sitem.space_after = Pt(20)

        self._add_footer(slide_summary, page_counter, 10, copyright_text)

        # Output logic
        if output_path:
            os.makedirs(os.path.dirname(output_path), exist_ok=True)
            prs.save(output_path)
            print(f"✅ Generated PowerPoint deck successfully: {output_path}")

        import io
        stream = io.BytesIO()
        prs.save(stream)
        return stream.getvalue()

pptx_generator_agent = PPTXGeneratorAgent()
