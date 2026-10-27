"""Cổng quiz lesson và quiz session (thư mục `Quizz lesson/` và `Quizz session/`).

Mỗi thư mục có đúng một tệp .md và đúng một tệp .xlsx (bỏ qua tệp khóa Office `~$*`).
`path` có thể là thư mục quiz hoặc một tệp nằm trong đó (khi đó dùng thư mục cha).

Cách xlsx mã hóa đáp án đúng (đo trên tệp thật của nodejs_ai_driven và reactjs_ai_driven):
cột `isCorrect` chứa chỉ số 1-4 (1 = A ... 4 = D), có thể là số nguyên hoặc chuỗi "1".."4".
Ở bộ lesson, các cột `right_1..right_6` luôn để trống; không dùng chúng để suy ra đáp án.

Tiêu đề thật của bộ session là `explanation_answer_1..4` (không lặp một tên); tiêu đề được so khớp chính xác
với tệp đã duyệt. Md session không có `**Giải thích:**`: lời giải từng phương án là dòng `- A: ...`.
"""
from __future__ import annotations

import re
import time
from dataclasses import dataclass
from pathlib import Path

from core.harness.findings import Finding, GateResult
from core.harness.gates import GATES
from core.harness.options import VerifyOptions

LESSON_QUESTIONS, SESSION_QUESTIONS = 5, 45
LESSON_TAGS = ("SYNTAX", "EXECUTION_FLOW", "CODE_TRACE", "COMPARISON", "TRAP_PREDICTION")
LESSON_SHEET, SESSION_SHEET = "Template Import Quiz", "Quiz Bank"
LESSON_HEADERS = (
    "question_content", "type", "points",
    "answer_1", "answer_2", "answer_3", "answer_4", "answer_5", "answer_6",
    "right_1", "right_2", "right_3", "right_4", "right_5", "right_6", "isCorrect",
)
SESSION_HEADERS = (
    "STT", "question_content",
    "answer_1", "explanation_answer_1", "answer_2", "explanation_answer_2",
    "answer_3", "explanation_answer_3", "answer_4", "explanation_answer_4",
    "isCorrect", "difficulty", "category",
)

_Q_RE = re.compile(r"^###\s+Câu\s+(\d+)\s*\((.+?)\)\s*$")
_OPT_RE = re.compile(r"^([A-D])\.\s+\S")
_ANS_RE = re.compile(r"\*\*Đáp án đúng:\*\*\s*([A-D])\b")
_FENCE = "```"


@dataclass
class QuizQuestion:
    number: int
    tag: str
    options: dict[str, str]
    answer: str
    has_explanation: bool
    has_feedback: bool


def parse_quiz_md(text: str) -> list[QuizQuestion]:
    """Tách các câu `### Câu N (TAG)`; phương án `A. `..`D. ` ngoài khối mã."""
    blocks: list[tuple[int, str, list[str]]] = []
    in_fence = False
    for line in text.splitlines():
        if line.lstrip().startswith(_FENCE):
            in_fence = not in_fence
        m = None if in_fence else _Q_RE.match(line)
        if m:
            blocks.append((int(m.group(1)), m.group(2).strip(), []))
        elif blocks:
            blocks[-1][2].append(line)
    questions = []
    for number, tag, lines in blocks:
        options: dict[str, str] = {}
        in_fence = False
        for line in lines:
            if line.lstrip().startswith(_FENCE):
                in_fence = not in_fence
                continue
            m = None if in_fence else _OPT_RE.match(line)
            if m and m.group(1) not in options:
                options[m.group(1)] = line[2:].strip()
        body = "\n".join(lines)
        am = _ANS_RE.search(body)
        questions.append(QuizQuestion(
            number, tag, options, am.group(1) if am else "",
            "**Giải thích:**" in body, "**Phản hồi khi chọn sai:**" in body,
        ))
    return questions


def _find_files(folder: Path, gate: str, result: GateResult) -> tuple[Path | None, Path | None]:
    found = {}
    for ext, label in ((".md", "md"), (".xlsx", "xlsx")):
        files = sorted(p for p in folder.iterdir()
                       if p.is_file() and p.suffix.lower() == ext and not p.name.startswith("~$"))
        if not files:
            result.findings.append(Finding("error", gate, f"Thư mục quiz thiếu tệp {label}", str(folder)))
        elif len(files) > 1:
            names = ", ".join(p.name for p in files)
            result.findings.append(Finding(
                "error", gate, f"Có nhiều tệp {label} ({names}), không xác định được tệp quiz", str(folder)))
        else:
            found[label] = files[0]
    return found.get("md"), found.get("xlsx")


def _read_rows(xlsx: Path, sheet: str, gate: str, result: GateResult) -> list[tuple] | None:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    except Exception as exc:  # tệp hỏng
        result.findings.append(Finding("error", gate, f"Không mở được xlsx: {exc}", xlsx.name))
        return None
    try:
        if sheet not in wb.sheetnames:
            result.findings.append(Finding(
                "error", gate, f"xlsx thiếu sheet `{sheet}` (có: {', '.join(wb.sheetnames)})", xlsx.name))
            return None
        return [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()


def _check_headers(rows: list[tuple], expected: tuple, gate: str, result: GateResult, name: str) -> bool:
    header = list(rows[0]) if rows else []
    while header and header[-1] is None:
        header.pop()
    if tuple(header) != expected:
        result.findings.append(Finding(
            "error", gate,
            f"Hàng tiêu đề xlsx không khớp: cần {len(expected)} tiêu đề ({', '.join(expected)}); "
            f"có {len(header)}: {', '.join(map(str, header))}",
            name))
        return False
    return True


def _data_rows(rows: list[tuple]) -> list[tuple]:
    return [r for r in rows[1:] if any(c is not None and str(c).strip() != "" for c in r)]


def _letter(value) -> str:
    try:
        n = int(str(value).strip())
    except (TypeError, ValueError):
        return ""
    return "ABCD"[n - 1] if 1 <= n <= 4 else ""


def _check_md(questions: list[QuizQuestion], expected: int, gate: str, result: GateResult, md_name: str,
              lesson: bool) -> None:
    def err(msg: str):
        result.findings.append(Finding("error", gate, msg, md_name))

    if len(questions) != expected:
        err(f"Quiz phải có đúng {expected} câu, hiện có {len(questions)}")
    numbers = [q.number for q in questions]
    if numbers != list(range(1, len(questions) + 1)):
        err(f"Số thứ tự câu không liên tục từ 1: {numbers}")
    for q in questions:
        label = f"Câu {q.number}"
        if lesson and q.tag not in LESSON_TAGS:
            err(f"{label}: nhãn `{q.tag}` không hợp lệ (cần một trong {', '.join(LESSON_TAGS)})")
        if list(q.options) != ["A", "B", "C", "D"]:
            err(f"{label}: cần đúng 4 phương án A-D, có {', '.join(q.options) or 'không có'}")
        if not q.answer:
            err(f"{label}: thiếu `**Đáp án đúng:**` hợp lệ (A-D)")
        if lesson:
            if not q.has_explanation:
                err(f"{label}: thiếu `**Giải thích:**`")
            if not q.has_feedback:
                err(f"{label}: thiếu `**Phản hồi khi chọn sai:**`")


def _check_quiz(path: Path, kind: str, lesson: bool) -> GateResult:
    start = time.perf_counter()
    path = Path(path)
    result = GateResult(kind=kind, path=str(path), gate=kind)
    gate = kind
    folder = path.parent if path.is_file() else path
    expected_n = LESSON_QUESTIONS if lesson else SESSION_QUESTIONS
    headers = LESSON_HEADERS if lesson else SESSION_HEADERS
    try:
        md, xlsx = _find_files(folder, gate, result)
        questions: list[QuizQuestion] = []
        if md is not None:
            try:
                questions = parse_quiz_md(md.read_text(encoding="utf-8"))
            except (OSError, UnicodeDecodeError) as exc:
                result.findings.append(Finding("error", gate, f"Không đọc được md: {exc}", md.name))
            else:
                _check_md(questions, expected_n, gate, result, md.name, lesson)
        if xlsx is not None:
            rows = _read_rows(xlsx, LESSON_SHEET if lesson else SESSION_SHEET, gate, result)
            if rows is not None and _check_headers(rows, headers, gate, result, xlsx.name):
                data = _data_rows(rows)
                if len(data) != expected_n:
                    result.findings.append(Finding(
                        "error", gate, f"xlsx phải có {expected_n} hàng dữ liệu, hiện có {len(data)}", xlsx.name))
                if md is not None and len(data) != len(questions):
                    result.findings.append(Finding(
                        "error", gate,
                        f"Số hàng dữ liệu xlsx ({len(data)}) khác số câu md ({len(questions)})", xlsx.name))
                col = headers.index("isCorrect")
                for row, q in zip(data, questions):
                    shown = row[col] if col < len(row) else None
                    got = _letter(shown)
                    if got != q.answer:
                        result.findings.append(Finding(
                            "error", gate,
                            f"Câu {q.number}: đáp án xlsx `isCorrect`={shown!r} ({got or 'không hợp lệ'}) "
                            f"khác đáp án md ({q.answer or 'thiếu'})",
                            xlsx.name))
        result.detail["questions"] = len(questions)
    except OSError as exc:
        result.findings.append(Finding("error", gate, f"Không đọc được thư mục quiz: {exc}", str(folder)))
    result.duration_s = time.perf_counter() - start
    return result


def check_quiz_lesson(path: Path, opts: VerifyOptions) -> GateResult:
    return _check_quiz(path, "quiz_lesson", lesson=True)


def check_quiz_session(path: Path, opts: VerifyOptions) -> GateResult:
    return _check_quiz(path, "quiz_session", lesson=False)


GATES["quiz_lesson"] = check_quiz_lesson
GATES["quiz_session"] = check_quiz_session
