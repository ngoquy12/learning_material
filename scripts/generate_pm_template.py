"""
scripts/generate_pm_template.py

Generates the Standard 10-Column PM Excel Template (`templates/PM_Template_Standard.xlsx`)
with Course Profile Header Metadata (Target Persona, CLOs, Capstone Target),
merged cells for Session/Type columns, beautiful OpenPyXL styling,
and sample data covering ORIENTATION, THEORY, PRACTICE, MINI_PROJECT, and FINAL_PROJECT sessions.
"""

import os
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_standard_pm_template():
    templates_dir = Path(__file__).resolve().parent.parent / "templates"
    templates_dir.mkdir(parents=True, exist_ok=True)
    output_path = templates_dir / "PM_Template_Standard.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PM_Syllabus"
    ws.views.sheetView[0].showGridLines = True

    # ── Course Profile Header Metadata (Rows 1..3) ───────────────────────────
    ws.append(["Chân Dung Sinh Viên Target (Persona)", "Sinh viên CNTT năm 1 hoặc người mới bắt đầu, chưa từng có tư duy lập trình trước đó."])
    ws.append(["Mục Tiêu Môn Học (Course Outcomes - CLO)", "Nắm vững cú pháp Python 3.x, tư duy thuật toán lặp/rẽ nhánh, lập trình OOP căn bản và xử lý dữ liệu File/SQLite."])
    ws.append(["Sản Phẩm Capstone Target", "Xây dựng Hệ thống Quản lý Bán hàng E-Commerce bằng Python Console kết hợp lưu trữ File JSON/SQLite."])
    ws.append([]) # Blank separator row (Row 4)

    # ── 10 Column Table Headers (Row 5) ──────────────────────────────────────
    headers = [
        "Session",
        "Loại Session",
        "Mã Session",
        "Tên Tiêu Đề Session",
        "Tên Lesson",
        "Nội Dung Chi Tiết (Lesson Scope)",
        "Kết Quả Mong Đợi (Expected Outcome)",
        "Phạm Vi CẤM DÙNG (Forbidden Scope)",
        "Phạm Vi ĐÃ HỌC (Allowed Scope)",
        "Tech Stack & Quy Chuẩn"
    ]

    ws.append(headers)

    # ── Sample Data Rows (Row 6 Onwards) ─────────────────────────────────────
    # Format: [Session, Loại Session, Mã Session, Tên Session, Tên Lesson, Nội Dung Chi Tiết, Expected Outcome, Forbidden, Allowed, Tech Stack]
    rows = [
        # Session 01 - ORIENTATION
        ["Session 01", "Orientation", "ORIENTATION", "Session 01 - Định hướng khóa học & Cài đặt môi trường", "Lesson 01: Giới thiệu lộ trình & Đồ án Capstone", "Định hướng môn học; Trình chiếu Demo Capstone Project; Quy định đánh giá và lộ trình học tập.", "Hiểu tổng quan môn học, lộ trình 30 buổi và tiêu chuẩn đánh giá Đồ án Capstone cuối khóa.", "CẤM: Code phức tạp, Algorithm, Loop, Function, Class.", "ĐÃ HỌC: Chưa có (Buổi mở đầu).", "Python 3.12, VS Code"],
        ["Session 01", "Orientation", "ORIENTATION", "Session 01 - Định hướng khóa học & Cài đặt môi trường", "Lesson 02: Hướng dẫn cài đặt môi trường lập trình", "Tải Python 3.x; Cấu hình PATH; Cài VS Code & Extensions (Python, Pylance); Tạo venv; Run Hello World.", "Cài đặt hoàn tất Python 3.12 & VS Code, tự chạy thành công file print('Hello World').", "CẤM: List, Dict, Loop, Exception, Function.", "ĐÃ HỌC: Tổng quan lộ trình môn học.", "Python 3.12, VS Code"],

        # Session 02 - THEORY
        ["Session 02", "Lý thuyết", "THEORY", "Session 02 - Biến, Kiểu dữ liệu & Nhập xuất dữ liệu", "Lesson 01: Khai báo Biến & Các kiểu dữ liệu cơ bản", "Khái niệm biến và ô nhớ; Biến int, float, str, bool; Quy tắc snake_case; Kiểm tra type() và id().", "Khai báo thành công các kiểu biến căn bản, phân biệt được int/float/str và áp dụng quy tắc đặt tên snake_case.", "CẤM: if-else, Loop, List, Dict, Function, Class.", "ĐÃ HỌC: Môi trường VS Code, Hello World.", "Python 3.12 (PEP 8)"],
        ["Session 02", "Lý thuyết", "THEORY", "Session 02 - Biến, Kiểu dữ liệu & Nhập xuất dữ liệu", "Lesson 02: Nhập xuất dữ liệu với print() và input()", "Hàm print() với sep, end; Hàm input() bàn phím trả về str; Ép kiểu int(), float(), str(); f-string.", "Viết được chương trình Console nhập xuất dữ liệu người dùng và in chuỗi kết quả chuẩn bằng f-string.", "CẤM: if-else, Loop, List, Dict, Function, Class.", "ĐÃ HỌC: Khai báo biến, int, float, str.", "Python 3.12 (PEP 8)"],

        # Session 03 - PRACTICE
        ["Session 03", "Thực hành", "PRACTICE", "Session 03 - Thực hành Biến, Kiểu dữ liệu & I/O", "Session 03 - Thực hành Biến, Kiểu dữ liệu & I/O", "Bộ 5 bài tập thực hành nhập xuất dữ liệu, tính toán biểu thức toán học cơ bản và ép kiểu f-string.", "Hoàn thành 5 bài tập tính toán biểu thức (tính chu vi/diện tích, tính hóa đơn) sử dụng input/print/f-string.", "CẤM: if-else, Loop, List, Dict, Function, Class.", "ĐÃ HỌC: Biến, int, float, str, print(), input().", "Python 3.12 (PEP 8)"],

        # Session 04 - THEORY
        ["Session 04", "Lý thuyết", "THEORY", "Session 04 - Toán tử & Cấu trúc điều kiện rẽ nhánh", "Lesson 01: Phép toán số học & Toán tử so sánh", "Phép toán +, -, *, /, //, %, **; Thứ tự PEMDAS; Các phép so sánh ==, !=, >, <, >=, <=; Boolean.", "Thực hiện thành công các phép tính chia lấy dư, lũy thừa và giải thích đúng biểu thức logic Boolean.", "CẤM: Loop, List, Dict, Function, Class.", "ĐÃ HỌC: Biến, kiểu dữ liệu, f-string.", "Python 3.12 (PEP 8)"],
        ["Session 04", "Lý thuyết", "THEORY", "Session 04 - Toán tử & Cấu trúc điều kiện rẽ nhánh", "Lesson 02: Cấu trúc rẽ nhánh if-elif-else & Match-case", "Cú pháp câu lệnh rẽ nhánh if, elif, else; Thụt lề Indentation 4 spaces; Match-case (Python 3.10+); Ternary.", "Viết được thuật toán kiểm tra rẽ nhánh 3-4 điều kiện phân loại và dùng match-case xử lý menu lệnh.", "CẤM: Loop, List, Dict, Function, Class.", "ĐÃ HỌC: Phép toán số học & so sánh.", "Python 3.12 (PEP 8)"],

        # Session 05 - PRACTICE
        ["Session 05", "Thực hành", "PRACTICE", "Session 05 - Thực hành Cấu trúc điều kiện rẽ nhánh", "Session 05 - Thực hành Cấu trúc điều kiện rẽ nhánh", "Bộ 5 bài tập thực hành rẽ nhánh phân cấp từ Dễ đến Xuất sắc: Kiểm tra số chẵn lẻ, Xếp loại học lực, Bậc giá điện.", "Giải quyết 100% 5 bài tập rẽ nhánh phân cấp phức tạp (tính tiền điện, xếp loại học lực).", "CẤM: Loop, List, Dict, Function, Class.", "ĐÃ HỌC: if-elif-else, match-case, ternary.", "Python 3.12 (PEP 8)"],

        # Session 06 - THEORY
        ["Session 06", "Lý thuyết", "THEORY", "Session 06 - Vòng lặp cơ bản & Thuật toán lặp", "Lesson 01: Vòng lặp for và hàm range()", "Cú pháp vòng lặp for; Hàm range(start, stop, step); Duyệt dải số; Tính tổng và đếm số.", "Viết được vòng lặp for duyệt dải số, áp dụng tính tổng chuỗi số và đếm số thỏa điều kiện.", "CẤM: List, Dict, Function, Class.", "ĐÃ HỌC: Biến, điều kiện rẽ nhánh if-else.", "Python 3.12 (PEP 8)"],
        ["Session 06", "Lý thuyết", "THEORY", "Session 06 - Vòng lặp cơ bản & Thuật toán lặp", "Lesson 02: Vòng lặp while & Điều khiển loop", "Cú pháp vòng lặp while theo điều kiện; break, continue, pass; Xử lý vòng lặp vô hạn.", "Xây dựng vòng lặp while lặp lại menu cho đến khi chọn thoát, điều khiển linh hoạt break/continue.", "CẤM: List, Dict, Function, Class.", "ĐÃ HỌC: Vòng lặp for, range().", "Python 3.12 (PEP 8)"],

        # Session 07 - PRACTICE
        ["Session 07", "Thực hành", "PRACTICE", "Session 07 - Thực hành Vòng lặp & Pattern Stars", "Session 07 - Thực hành Vòng lặp & Pattern Stars", "Bộ 5 bài tập thực hành vòng lặp for/while, vẽ các hình pattern dấu sao (*), tính tổng chuỗi số.", "Vẽ được các hình sao tam giác/hình vuông bằng vòng lặp lồng nhau và giải bài toán tính tổng số học.", "CẤM: List, Dict, Function, Class.", "ĐÃ HỌC: for, while, range(), break, continue.", "Python 3.12 (PEP 8)"],

        # Session 08 - MINI_PROJECT
        ["Session 08", "Mini Project", "MINI_PROJECT", "Session 08 - Mini Project 1: Tính toán cước vận chuyển tự động", "Session 08 - Mini Project 1: Tính toán cước vận chuyển tự động", "4 Entry Tests; Tài liệu SRS spec tinh gọn; Đề bài Mini Project tính cước vận chuyển E-Commerce theo bậc giá.", "Hoàn thành Mini Project 1 chạy Console hoàn chỉnh, vượt qua 100% 4 Entry Test cases quy định.", "CẤM: List, Dict, Function, Class.", "ĐÃ HỌC: Toàn bộ kiến thức Session 01 - 07.", "Python 3.12 (PEP 8)"],

        # Session 30 - FINAL_PROJECT
        ["Session 30", "Dự án cuối khóa", "FINAL_PROJECT", "Session 30 - Thi kết thúc môn Python Core (Capstone Project)", "Session 30 - Thi kết thúc môn Python Core (Capstone Project)", "Đồ án Capstone cuối khóa; Sơ đồ Kiến trúc CSDL; Tài liệu đặc tả SRS đầy đủ; Rubric chấm Đồ án 100đ.", "Thuyết trình và bảo vệ thành công Đồ án Capstone Python Console đáp ứng 100% Rubric đánh giá.", "CẤM: Không giới hạn.", "ĐÃ HỌC: Toàn bộ kiến thức toàn khóa học.", "Python 3.12 (PEP 8)"]
    ]

    for row in rows:
        ws.append(row)

    # Merge Cells Range for Sessions with multiple lessons (Header is now row 5, data starts row 6)
    # Session 01: rows 6..7
    # Session 02: rows 8..9
    # Session 04: rows 11..12
    # Session 06: rows 14..15
    merge_ranges = [
        (6, 7),   # Session 01
        (8, 9),   # Session 02
        (11, 12), # Session 04
        (14, 15)  # Session 06
    ]

    for start_r, end_r in merge_ranges:
        ws.merge_cells(start_row=start_r, end_row=end_r, start_column=1, end_column=1) # Col A: Session
        ws.merge_cells(start_row=start_r, end_row=end_r, start_column=2, end_column=2) # Col B: Loại Session
        ws.merge_cells(start_row=start_r, end_row=end_r, start_column=3, end_column=3) # Col C: Mã Session
        ws.merge_cells(start_row=start_r, end_row=end_r, start_column=4, end_column=4) # Col D: Tên Session

    # ── Styling ─────────────────────────────────────────────────────────────
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Navy Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    meta_label_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid") # Soft Blue
    meta_label_font = Font(name="Calibri", size=10, bold=True, color="1F4E78")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    # Apply Metadata Header Styling (Rows 1..3)
    for r_idx in range(1, 4):
        ws.row_dimensions[r_idx].height = 22
        cell_lbl = ws.cell(row=r_idx, column=1)
        cell_lbl.fill = meta_label_fill
        cell_lbl.font = meta_label_font
        cell_lbl.alignment = Alignment(horizontal="left", vertical="center")
        cell_lbl.border = thin_border
        
        # Merge metadata value across columns B..J
        ws.merge_cells(start_row=r_idx, end_row=r_idx, start_column=2, end_column=10)
        cell_val = ws.cell(row=r_idx, column=2)
        cell_val.font = Font(name="Calibri", size=10, italic=True)
        cell_val.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell_val.border = thin_border

    # Apply 10-Column Table Header Styles (Row 5)
    for col_idx in range(1, 11):
        cell = ws.cell(row=5, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[5].height = 28

    # Apply Data Cell Styles (Row 6 Onwards)
    for row_idx in range(6, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 24
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, 11):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            
            if is_even and not cell.fill.fill_type:
                cell.fill = zebra_fill
                
            # Alignment rules
            if col_idx in [1, 2, 3]:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif col_idx in [4, 5]:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Set Column Widths (10 Columns)
    col_widths = [14, 16, 16, 36, 36, 45, 45, 35, 35, 22]
    for col_idx, width in enumerate(col_widths, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    wb.save(output_path)
    print(f"[PM Template Generator] Successfully generated 10-Column Standard PM Template: {output_path}")
    return output_path

if __name__ == "__main__":
    generate_standard_pm_template()
