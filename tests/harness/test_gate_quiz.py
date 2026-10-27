"""Cổng quiz lesson và quiz session.

Cách xlsx mã hóa đáp án đúng (đo trên tệp thật của nodejs_ai_driven/reactjs_ai_driven):
cột `isCorrect` chứa chỉ số 1-4 (1 = A ... 4 = D), có thể là số nguyên hoặc chuỗi;
các cột `right_1..right_6` của bộ lesson luôn để trống.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from core.harness.gates import GATES
from core.harness.gates.quiz import (
    LESSON_HEADERS, SESSION_HEADERS, check_quiz_lesson, check_quiz_session, parse_quiz_md,
)
from core.harness.options import VerifyOptions

OPTS = VerifyOptions()
TAGS = ["SYNTAX", "EXECUTION_FLOW", "CODE_TRACE", "COMPARISON", "TRAP_PREDICTION"]
ANS = "ABCDA"


def lesson_md(n=5, tags=None, drop_feedback_at=None, answers=ANS) -> str:
    tags = tags or TAGS
    out = ["# Quiz - Demo", ""]
    for i in range(1, n + 1):
        out += [f"### Câu {i} ({tags[(i - 1) % len(tags)]})", f"Nội dung câu {i}?", "",
                "A. một", "B. hai", "C. ba", "D. bốn", "",
                f"**Đáp án đúng:** {answers[(i - 1) % len(answers)]}", "", "**Giải thích:** vì lý do.", ""]
        if i != drop_feedback_at:
            out += ["**Phản hồi khi chọn sai:** A sai vì.", ""]
        out += ["---", ""]
    return "\n".join(out)


def session_md(n=45) -> str:
    out = ["# Quizz cuối giờ", "", "## Nhóm: Ứng dụng thực tế (độ khó 6)", ""]
    for i in range(1, n + 1):
        out += [f"### Câu {i} (Lesson {i % 7 + 1:02d})", f"Câu hỏi {i}?", "",
                "A. một", "B. hai", "C. ba", "D. bốn", "", f"**Đáp án đúng:** {'ABCD'[i % 4]}", "",
                "- A: sai vì.", "- B: sai vì.", "- C: sai vì.", "- D: sai vì.", "", "---", ""]
    return "\n".join(out)


def lesson_xlsx(path: Path, rows=5, answers=(1, 2, 3, 4, 1), headers=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template Import Quiz"
    ws.append(headers or LESSON_HEADERS)
    for i in range(rows):
        ws.append([f"<p>q{i}</p>", "SINGLE_CHOICE", 1, "a", "b", "c", "d", None, None,
                   None, None, None, None, None, None, str(answers[i % len(answers)])])
    wb.save(path)


def session_xlsx(path: Path, rows=45, headers=None, shift=0):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quiz Bank"
    ws.append(headers or SESSION_HEADERS)
    for i in range(1, rows + 1):
        ws.append([i, f"q{i}", "a", "ea", "b", "eb", "c", "ec", "d", "ed", (i + shift) % 4 + 1, 6, "BÀI MỚI"])
    wb.save(path)


def make_lesson(tmp_path, md=None, **kw) -> Path:
    d = tmp_path / "Lesson 01 - X" / "Quizz lesson"
    d.mkdir(parents=True)
    (d / "quiz.md").write_text(md if md is not None else lesson_md(), encoding="utf-8")
    lesson_xlsx(d / "Quizz_S01_L01.xlsx", **kw)
    return d


def make_session(tmp_path, md=None, **kw) -> Path:
    d = tmp_path / "Quizz session"
    d.mkdir(parents=True)
    (d / "quizz_cuoi_gio.md").write_text(md if md is not None else session_md(), encoding="utf-8")
    session_xlsx(d / "Session01.xlsx", **kw)
    return d


def errs(r):
    return [f.message for f in r.findings if f.severity == "error"]


def test_registered():
    assert GATES["quiz_lesson"] is check_quiz_lesson
    assert GATES["quiz_session"] is check_quiz_session


def test_parse_quiz_md():
    qs = parse_quiz_md(lesson_md())
    assert [q.number for q in qs] == [1, 2, 3, 4, 5]
    assert qs[0].tag == "SYNTAX" and qs[0].answer == "A" and qs[1].answer == "B"
    assert list(qs[0].options) == ["A", "B", "C", "D"]
    assert qs[0].has_explanation and qs[0].has_feedback
    assert not parse_quiz_md(lesson_md(drop_feedback_at=3))[2].has_feedback


def test_lesson_ok(tmp_path):
    r = check_quiz_lesson(make_lesson(tmp_path), OPTS)
    assert r.status == "pass", errs(r)
    assert r.kind == "quiz_lesson"


def test_lesson_path_may_be_file(tmp_path):
    d = make_lesson(tmp_path)
    assert check_quiz_lesson(d / "quiz.md", OPTS).status == "pass"
    assert check_quiz_lesson(d / "Quizz_S01_L01.xlsx", OPTS).status == "pass"


def test_lesson_missing_feedback(tmp_path):
    r = check_quiz_lesson(make_lesson(tmp_path, md=lesson_md(drop_feedback_at=3)), OPTS)
    assert r.status == "fail" and any("Câu 3" in m for m in errs(r))


def test_lesson_bad_tag(tmp_path):
    tags = ["SYNTAXX"] + TAGS[1:]
    r = check_quiz_lesson(make_lesson(tmp_path, md=lesson_md(tags=tags)), OPTS)
    assert any("SYNTAXX" in m for m in errs(r))


def test_lesson_wrong_question_count(tmp_path):
    r = check_quiz_lesson(make_lesson(tmp_path, md=lesson_md(n=4)), OPTS)
    assert any("5" in m and "4" in m for m in errs(r))


def test_lesson_xlsx_row_count(tmp_path):
    r = check_quiz_lesson(make_lesson(tmp_path, rows=4), OPTS)
    assert any("hàng dữ liệu" in m for m in errs(r))


def test_lesson_xlsx_answer_mismatch(tmp_path):
    r = check_quiz_lesson(make_lesson(tmp_path, answers=(1, 3, 3, 4, 1)), OPTS)
    msgs = errs(r)
    assert any("Câu 2" in m for m in msgs) and not any("Câu 3" in m for m in msgs)


def test_lesson_xlsx_header_exact(tmp_path):
    h = list(LESSON_HEADERS)
    h[5] = "answer_3 "
    r = check_quiz_lesson(make_lesson(tmp_path, headers=h), OPTS)
    assert any("tiêu đề" in m for m in errs(r))


def test_missing_xlsx_or_md(tmp_path):
    d = make_lesson(tmp_path)
    next(d.glob("*.xlsx")).unlink()
    assert any("xlsx" in m for m in errs(check_quiz_lesson(d, OPTS)))
    (d / "quiz.md").unlink()
    assert any("md" in m for m in errs(check_quiz_lesson(d, OPTS)))


def test_ambiguous_and_lock_files(tmp_path):
    d = make_lesson(tmp_path)
    (d / "~$Quizz_S01_L01.xlsx").write_bytes(b"lock")
    assert check_quiz_lesson(d, OPTS).status == "pass"
    (d / "other.md").write_text("x", encoding="utf-8")
    assert any("nhiều" in m for m in errs(check_quiz_lesson(d, OPTS)))


def test_session_ok(tmp_path):
    r = check_quiz_session(make_session(tmp_path), OPTS)
    assert r.status == "pass", errs(r)
    assert r.kind == "quiz_session"


def test_session_44_questions(tmp_path):
    r = check_quiz_session(make_session(tmp_path, md=session_md(44)), OPTS)
    assert any("45" in m for m in errs(r))


def test_session_xlsx_rows_and_mismatch(tmp_path):
    assert any("hàng dữ liệu" in m for m in errs(check_quiz_session(make_session(tmp_path, rows=44), OPTS)))
    r = check_quiz_session(make_session(tmp_path / "b", shift=1), OPTS)
    assert any("Câu 1" in m for m in errs(r))


def test_session_header_exact(tmp_path):
    h = list(SESSION_HEADERS)
    h[3] = "explanation_answer"
    r = check_quiz_session(make_session(tmp_path, headers=h), OPTS)
    assert any("tiêu đề" in m for m in errs(r))


def test_session_does_not_require_explanation_block(tmp_path):
    assert "Giải thích" not in session_md()
    assert check_quiz_session(make_session(tmp_path), OPTS).status == "pass"


def _two_sets(tmp_path, extra=None) -> Path:
    d = tmp_path / "Quizz session"
    d.mkdir(parents=True)
    for tok in ("dau_gio", "cuoi_gio"):
        (d / f"quizz_{tok}.md").write_text(session_md(), encoding="utf-8")
        session_xlsx(d / f"Session01._Quizz_{tok.title()}_X.xlsx")
    for name, content in (extra or {}).items():
        (d / name).write_text(content, encoding="utf-8")
    return d


def test_session_two_sets_pass_and_json_ignored(tmp_path):
    d = _two_sets(tmp_path, {"quizz_dau_gio.json": "{}", "quizz_cuoi_gio.json": "{}"})
    r = check_quiz_session(d, OPTS)
    assert r.status == "pass", errs(r)
    assert set(r.detail["questions"]) == {"dau_gio", "cuoi_gio"}


def test_session_unpaired_md(tmp_path):
    d = _two_sets(tmp_path)
    next(d.glob("*Dau_Gio*.xlsx")).unlink()
    msgs = errs(check_quiz_session(d, OPTS))
    assert any("thiếu tệp xlsx" in m and "dau_gio" in m for m in msgs)


def test_session_two_md_same_token(tmp_path):
    d = _two_sets(tmp_path, {"quizz_cuoi_gio_v2.md": session_md()})
    assert any("nhiều tệp md" in m and "cuoi_gio" in m for m in errs(check_quiz_session(d, OPTS)))


def test_session_error_names_set(tmp_path):
    d = _two_sets(tmp_path)
    (d / "quizz_dau_gio.md").write_text(session_md(44), encoding="utf-8")
    assert any(m.startswith("[bộ dau_gio]") and "45" in m for m in errs(check_quiz_session(d, OPTS)))


def test_md_numbering_not_sequential(tmp_path):
    md = lesson_md().replace("### Câu 3 ", "### Câu 7 ")
    assert any("liên tục" in m for m in errs(check_quiz_lesson(make_lesson(tmp_path, md=md), OPTS)))


def test_lesson_missing_explanation(tmp_path):
    md = lesson_md().replace("**Giải thích:** vì lý do.", "", 1)
    assert any("Câu 1" in m and "Giải thích" in m for m in errs(check_quiz_lesson(make_lesson(tmp_path, md=md), OPTS)))


def test_fewer_than_four_options(tmp_path):
    md = lesson_md().replace("D. bốn", "", 1)
    assert any("Câu 1" in m and "4 phương án" in m for m in errs(check_quiz_lesson(make_lesson(tmp_path, md=md), OPTS)))


def test_missing_answer(tmp_path):
    md = lesson_md().replace("**Đáp án đúng:** A", "", 1)
    assert any("Câu 1" in m and "Đáp án" in m for m in errs(check_quiz_lesson(make_lesson(tmp_path, md=md), OPTS)))


def test_int_iscorrect_accepted_and_non_numeric_rejected(tmp_path):
    d = make_lesson(tmp_path)
    x = next(d.glob("*.xlsx"))
    wb = openpyxl.load_workbook(x)
    ws = wb.active
    ws.cell(2, 16).value = 1  # int, đúng
    ws.cell(3, 16).value = "B?"  # không phải số
    wb.save(x)
    msgs = errs(check_quiz_lesson(d, OPTS))
    assert len(msgs) == 1 and "Câu 2" in msgs[0] and "không hợp lệ" in msgs[0]


def test_blank_row_in_middle_flagged(tmp_path):
    d = make_lesson(tmp_path)
    x = next(d.glob("*.xlsx"))
    wb = openpyxl.load_workbook(x)
    wb.active.insert_rows(4)
    wb.save(x)
    assert any("hàng trống" in m for m in errs(check_quiz_lesson(d, OPTS)))


def test_session_heading_without_tag_accepted(tmp_path):
    md = re.sub(r"^(### Câu \d+) \(Lesson \d+\)$", r"\1", session_md(), flags=re.M)
    assert "(Lesson" not in md
    r = check_quiz_session(make_session(tmp_path, md=md), OPTS)
    assert r.status == "pass", errs(r)


def test_lesson_heading_without_tag_still_fails(tmp_path):
    md = re.sub(r"^(### Câu \d+) \(\w+\)$", r"\1", lesson_md(), flags=re.M)
    r = check_quiz_lesson(make_lesson(tmp_path, md=md), OPTS)
    assert any("nhãn" in m for m in errs(r))


def test_no_question_heading_clear_error(tmp_path):
    md = session_md().replace("### Câu", "#### Cau")
    r = check_quiz_session(make_session(tmp_path, md=md), OPTS)
    assert any("không nhận ra tiêu đề câu hỏi dạng '### Câu N'" in m for m in errs(r)), errs(r)
