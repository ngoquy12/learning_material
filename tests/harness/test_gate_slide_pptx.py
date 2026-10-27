"""Cổng slide pptx: kiểm tĩnh bằng check_deck, lint slides.json, render tùy chọn."""
from __future__ import annotations

import json
import sys

import pytest

pytest.importorskip("pptx")
from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from core.harness.gates import GATES  # noqa: E402
from core.harness.gates.slide_pptx import check_slide_pptx  # noqa: E402
from core.harness.options import VerifyOptions  # noqa: E402


def make_deck(path, n=2, size_pt=18):
    prs = Presentation()
    for i in range(n):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = f"Nội dung slide {i + 1}"
        run.font.size = Pt(size_pt)
        s.notes_slide.notes_text_frame.text = "Ghi chú giảng"
    prs.save(str(path))


def folder(tmp_path):
    d = tmp_path / "Slide bai giang"
    d.mkdir()
    return d


def test_registered():
    assert GATES["slide_pptx"] is check_slide_pptx


def test_clean_deck_passes(tmp_path):
    p = folder(tmp_path) / "deck.pptx"
    make_deck(p)
    r = check_slide_pptx(p, VerifyOptions())
    assert r.status == "pass", r.findings
    assert r.detail["slides"] == 2


def test_static_error_mapped(tmp_path):
    p = folder(tmp_path) / "deck.pptx"
    make_deck(p, size_pt=8)
    r = check_slide_pptx(p, VerifyOptions())
    assert r.status == "fail"
    assert any(f.severity == "error" and "cỡ chữ" in f.message for f in r.findings)


def test_slides_json_overflow_is_error(tmp_path):
    d = folder(tmp_path)
    p = d / "deck.pptx"
    make_deck(p)
    slide = {"type": "comparison_2col", "h2": "So sánh", "left_title": "A", "right_title": "B",
             "left_items": [f"mục {i}" for i in range(6)], "right_items": ["x"]}
    (d / "slides.json").write_text(json.dumps({"slides": [slide]}, ensure_ascii=False), encoding="utf-8")
    r = check_slide_pptx(p, VerifyOptions())
    assert any(f.severity == "error" and "slide 1" in f.message for f in r.findings), r.findings


def test_unreadable_pptx_is_error(tmp_path):
    p = folder(tmp_path) / "deck.pptx"
    p.write_bytes(b"not a zip")
    r = check_slide_pptx(p, VerifyOptions())
    assert r.status == "fail"


def test_render_skipped_off_windows(tmp_path, monkeypatch):
    p = folder(tmp_path) / "deck.pptx"
    make_deck(p, size_pt=8)
    monkeypatch.setattr(sys, "platform", "linux")
    r = check_slide_pptx(p, VerifyOptions(render=True))
    assert r.skipped_reason
    assert any(f.severity == "error" for f in r.findings)  # lỗi tĩnh vẫn giữ
